import os
import io
import time
import requests
import streamlit as st
import pandas as pd
from datetime import datetime
from openpyxl.styles import Border, Side

st.set_page_config(page_title="Magazzino - WinGaming", page_icon="📦", layout="wide")

# 🎨 STILE GRAFICO DI MANUELA: ANNIENTAMENTO RIGIDO E TOTALE DEL ROSSO/ARANCIONE DALLE CAPSULE
st.markdown("""
    <style>
    /* 🎯 SOVRASCRITTURA DI TUTTE LEGATE VARIABILI INTERNE DI STREAMLIT (ELIMINATO IL ROSSO OVUNQUE) */
    :root {
        --primary-color: #e2e8f0 !important;
        --colors-primary: #e2e8f0 !important;
        --colors-primary-text: #334155 !important;
    }
    
    #MainMenu, footer, .stDecoration, [data-testid="stFooter"] { visibility: hidden !important; display: none !important; }
    .stApp { background-color: #f8fafc !important; color: #1e293b !important; font-family: 'Segoe UI', sans-serif !important; }
    
    h1, [data-testid="stHeader"] h1 { color: #0f766e !important; font-size: 20px !important; text-align: center; font-weight: 700 !important; margin-bottom: 10px; }
    h3, .stMarkdown h3 { color: #1e293b !important; font-size: 15px !important; font-weight: 600 !important; margin-top: 10px !important; margin-bottom: 5px !important; }
    
    /* Pannello badge informativo grigio-azzurro delicato */
    .time-badge { 
        background-color: #f1f5f9 !important; 
        border-left: 5px solid #0f766e !important; 
        padding: 10px !important; 
        border-radius: 6px !important; 
        font-size: 13px !important; 
        color: #475569 !important; 
        font-weight: 500;
        margin-bottom: 15px !important;
    }
    
    /* Bottoni contabili chiari */
    .stButton > button, .stDownloadButton > button { 
        background-color: #f8fafc !important; color: #334155 !important; border: 1px solid #e2e8f0 !important; 
        font-weight: 400 !important; font-size: 13px !important; width: 100% !important; border-radius: 8px !important; height: 36px !important; 
    }
    .stButton > button:hover, .stDownloadButton > button:hover { background-color: #f1f5f9 !important; border-color: #cbd5e1 !important; color: #0f172a !important; }
    
    /* 🎯 PIALLATURA TOTALE DEL ROSSO DAL CONTENITORE DELLE PILLOLE SCELTE */
    div[data-testid="stMultiSelectFloatingInlineValue"] > div,
    div[data-testid="stMultiSelectFloatingInlineValue"] div,
    span[data-testid="stTag"],
    span[data-testid="stTag"] > span,
    [data-baseweb="tag"],
    [data-baseweb="tag"] span,
    div[role="button"][tabindex="0"] {
        background-color: #e2e8f0 !important; /* Sfondo grigio-azzurro chiaro pastello */
        color: #334155 !important;            /* Testo grigio scuro contabile */
        border: 1px solid #cbd5e1 !important;  /* Bordino delicato */
        border-radius: 4px !important;
    }
    
    /* Impedisce al testo interno di riaccendersi in arancione o bianco */
    div[data-testid="stMultiSelectFloatingInlineValue"] span,
    span[data-testid="stTag"] span,
    [data-baseweb="tag"] span,
    div[role="button"] span {
        color: #334155 !important;
        font-size: 13px !important;
        font-weight: 500 !important;
    }
    
    /* Sforza le icone "X" di cancellazione a rimanere grigio scuro senza arrossare */
    div[data-testid="stMultiSelectFloatingInlineValue"] svg,
    span[data-testid="stTag"] svg,
    [data-baseweb="tag"] svg,
    div[role="button"] svg {
        fill: #475569 !important;
        color: #475569 !important;
    }
    </style>
""", unsafe_allow_html=True)

# 🛡️ PRIVACY RUOLO UTENTE
ruolo_utente = str(st.session_state.get("user_ruolo", "TECNICO")).strip().upper()
if ruolo_utente not in ["ADMIN", "SUPERVISORE", "UFFICIO"]:
    st.error("🔒 Accesso Riservato alla direzione.")
    st.stop()

# 📂 CONFIGURAZIONE ROTTE FILE INTERNE
FILE_GIACENZA_SCHEDE = "giacenza_schede.xlsx"
FILE_SLOT_PORTALI = "elenco_slot_portali.xlsx"
FILE_ELENCO_MAGAZZINI = "elenco_magazzini.xlsx"

# Titoli ripristinati a dimensions eleganti di fabbrica WinGaming
st.markdown("<h1>📦 Hub Magazzino — Estrazione a Comando Sansone</h1>", unsafe_allow_html=True)
st.markdown("<h3>🔄 Sincronizzazione Registri in Tempo Reale</h3>", unsafe_allow_html=True)

# Recupero immediato del token di sicurezza dai Secrets aziendali (Chiave nativa di Manuela)
t_git = str(st.secrets["github"]["token_accesso"]).strip()
s_api = "api" + "." + "github" + "." + "com"
repo_path = "wingamingsrl/AppWGt"
workflow_file = "cron_magazzino_manuale.yml"

if st.button("🚀 AVVIA ESTRAZIONE REALE DA SANSONE", key="btn_lancio_magazzino_manuale"):
    try:
        url_wf_magazzino = f"https://{s_api}/repos/{repo_path}/actions/workflows/{workflow_file}/dispatches"
        headers_lodi = {"Authorization": f"token {t_git}", "Accept": "application/vnd.github+json", "User-Agent": "WinGaming-Cloud-App"}
        res_lodi = requests.post(url_wf_magazzino, json={"ref": "main"}, headers=headers_lodi, timeout=10)
        
        if res_lodi.status_code == 204 or res_lodi.status_code == 202:
            # 🎯 IL RADAR SPIA DI MANUELA: Intercetta la fine reale dell'Action senza usare i secondi!
            stato_attesa = st.empty()
            stato_attesa.info("⏳ Robot Sansone avviato nel Cloud... Sto piantonando l'Action su GitHub... Non toccare nulla.")
            
            time.sleep(6.0) # Pausa tecnica per dare tempo a GitHub di registrare e far partire la nuova esecuzione
            
            url_runs_check = f"https://{s_api}/repos/{repo_path}/actions/workflows/{workflow_file}/runs?per_page=1"
            
            completato = False
            tentativi = 0
            max_tentativi = 40 # Protezione paracadute: massimo 3 minuti e mezzo di ascolto totale
            
            while not completato and tentativi < max_tentativi:
                tentativi += 1
                try:
                    res_check = requests.get(url_runs_check, headers=headers_lodi, timeout=5)
                    if res_check.status_code == 200:
                        dati_runs = res_check.json()
                        if dati_runs.get("workflow_runs"):
                            run_corrente = dati_runs["workflow_runs"][0]
                            status_action = run_corrente.get("status", "").strip().lower()
                            conclusion_action = run_corrente.get("conclusion", "")
                            
                            # Se lo stato non è più "in_progress" o "queued", significa che ha FINITO!
                            if status_action == "completed":
                                completato = True
                                if str(conclusion_action).strip().lower() == "success":
                                    st.cache_data.clear()
                                    stato_attesa.success("✅ AGGIORNAMENTO COMPLETATO! Il robot ha scritto i file Excel e la plancia viene rinfrescata!")
                                else:
                                    stato_attesa.error("❌ ERRORE CRITICO: L'Action è terminata ma il robot ha segnalato un'anomalia nei registri di Sansone.")
                            else:
                                # Calcola un contatore visivo per l'utente in ufficio
                                secondi_stimati = tentativi * 5
                                stato_attesa.info(f"⚙️ Il server cloud sta elaborando la giacenza... (Fase di calcolo: {secondi_stimati}s). Aspetto che finisca l'Action...")
                    
                    if not completato:
                        time.sleep(5.0) # Interroga il server di GitHub ogni 5 secondi spaccati
                except Exception:
                    time.sleep(5.0)
            
            if not completato:
                st.warning("⏳ Il server cloud sta impiegando più tempo del previsto. Forza il rinfresco manuale tra qualche istante.")
            
            time.sleep(1.5)
            st.rerun()
        else:
            st.error(f"❌ Impossibile avviare il robot. Risposta server: {res_lodi.status_code} - {res_lodi.text}")
            time.sleep(4.0)
    except Exception as e_lodi_click:
        st.error(f"💥 Errore di rete interno: {str(e_lodi_click)}")
        time.sleep(4.0)
# 🎯 LA DOPPIA BARRIERA DI MANUELA: Verifica che i file esistano e non siano corrotti o vuoti (peso > 10KB)
file_realmente_valido = False
if os.path.exists(FILE_GIACENZA_SCHEDE) and os.path.getsize(FILE_GIACENZA_SCHEDE) > 10000:
    file_realmente_valido = True

data_ora_aggiornamento = datetime.now().strftime("%d/%m/%Y alle ore %H:%M:%S")

# Se il file supera il controllo di peso, interroga le API di GitHub per mostrare la data di completamento reale
if file_realmente_valido:
    try:
        s_api = "api" + "." + "github" + "." + "com"
        url_runs = f"https://{s_api}/repos/{repo_path}/actions/workflows/{workflow_file}/runs?status=success&per_page=1"
        headers_runs = {"Authorization": f"token {t_git}", "Accept": "application/vnd.github+json", "User-Agent": "WinGaming-Cloud-App"}
        res_runs = requests.get(url_runs, headers=headers_runs, timeout=5)
        if res_runs.status_code == 200:
            dati_runs = res_runs.json()
            if dati_runs.get("workflow_runs"):
                conclusa_at = dati_runs["workflow_runs"][0]["updated_at"]
                dt_utc = datetime.strptime(conclusa_at, "%Y-%m-%dT%H:%M:%SZ")
                import datetime as dt_mod
                dt_locale = dt_utc + dt_mod.timedelta(hours=2)
                data_ora_aggiornamento = dt_locale.strftime("%d/%m/%Y alle ore %H:%M:%S")
    except Exception:
        # Paracadute: se le API di GitHub fossero temporaneamente offline, legge la data fisica del file locale
        data_ora_aggiornamento = datetime.fromtimestamp(os.path.getmtime(FILE_GIACENZA_SCHEDE)).strftime("%d/%m/%Y alle ore %H:%M:%S")

# Visualizzazione del badge temporale protetto
if file_realmente_valido:
    st.markdown(f"<div class='time-badge'>📅 <b>Ultimo aggiornamento reale della giacenza (GitHub Actions):</b> {data_ora_aggiornamento}</div>", unsafe_allow_html=True)
else:
    st.markdown("<div class='time-badge'>📅 <b>Stato registro giacenze:</b> ⚠️ Attenzione: il file Excel è mancante o corrotto dall'ultimo scaricamento. Sincronizza nuovamente.</div>", unsafe_allow_html=True)

# Ripresa dei caricamenti Excel standard del magazzino
if file_realmente_valido and os.path.exists(FILE_GIACENZA_SCHEDE):
    try:
        df_schede = pd.read_excel(FILE_GIACENZA_SCHEDE).fillna("")
        df_schede.columns = [str(c).strip() for c in df_schede.columns]





        # 🛡️ ESTRAZIONE OPZIONI TENDINA DIRETTAMENTE DAL FILE ANAGRAFICA MAGAZZINI (HEADER=1)
        opzioni_totali_tendina = []
        dizionario_magazzini = {}
        
        if os.path.exists(FILE_ELENCO_MAGAZZINI):
            df_magazzini_reali = pd.read_excel(FILE_ELENCO_MAGAZZINI, header=1).fillna("")
            df_magazzini_reali.columns = [str(c).strip() for c in df_magazzini_reali.columns]
            
            col_id_mag = "ID" if "ID" in df_magazzini_reali.columns else (str(df_magazzini_reali.columns) if len(df_magazzini_reali.columns) > 0 else "")
            col_nome_mag = "Nome" if "Nome" in df_magazzini_reali.columns else ("Descrizione" if "Descrizione" in df_magazzini_reali.columns else (str(df_magazzini_reali.columns) if len(df_magazzini_reali.columns) > 1 else col_id_mag))
            
            df_magazzini_reali[col_id_mag] = df_magazzini_reali[col_id_mag].astype(str).apply(lambda x: str(x).strip())
            df_magazzini_reali[col_nome_mag] = df_magazzini_reali[col_nome_mag].astype(str).apply(lambda x: str(x).strip())
            dizionario_magazzini = dict(zip(df_magazzini_reali[col_id_mag], df_magazzini_reali[col_nome_mag]))
            
            greggio_magazzini = df_magazzini_reali[col_nome_mag].drop_duplicates().tolist()
            opzioni_totali_tendina = sorted([str(m).strip() for m in greggio_magazzini if str(m).strip().upper() not in ["NOME", "DESCRIZIONE", "ID", "CODICE", ""]])

        col_luogo_schede = "Luogo" if "Luogo" in df_schede.columns else df_schede.columns[-1]
        df_schede[col_luogo_schede] = df_schede[col_luogo_schede].astype(str).apply(lambda x: str(x).strip())
        df_schede[col_luogo_schede] = df_schede[col_luogo_schede].map(dizionario_magazzini).fillna(df_schede[col_luogo_schede])
        
        if not opzioni_totali_tendina:
            tutti_i_luoghi_reali = df_schede[col_luogo_schede].drop_duplicates().tolist()
            opzioni_totali_tendina = sorted([str(l).strip() for l in tutti_i_luoghi_reali if str(l).strip().upper() not in ["LUOGO", "MAGAZZINO SCONOSCIUTO", ""]])
        
        # 🎯 DIZIONARIO ALFANUMERICO COMPLETO PER IL CERCA.VERT SUL CODICE AAMS (CON HEADER=1)
        mappa_giorni_decaden = {}
        if os.path.exists(FILE_SLOT_PORTALI):
            df_slot = pd.read_excel(FILE_SLOT_PORTALI, header=1).fillna("")
            df_slot.columns = [str(c).strip() for c in df_slot.columns]
            
            col_id_slot = "Codice AAMS" if "Codice AAMS" in df_slot.columns else ("Codice Aams" if "Codice Aams" in df_slot.columns else str(df_slot.columns))
            col_giorni_origine = "Giorni Decadenza" if "Giorni Decadenza" in df_slot.columns else str(df_slot.columns[-1])
            
            for idx, riga_slot in df_slot.iterrows():
                val_id_slot = str(riga_slot[col_id_slot]).strip()
                if val_id_slot.endswith('.0'): val_id_slot = val_id_slot[:-2]
                val_giorni = str(riga_slot[col_giorni_origine]).strip()
                if val_giorni.endswith('.0'): val_giorni = val_giorni[:-2]
                if val_id_slot and val_id_slot not in ["nan", "None", "0", "0.0"]:
                    mappa_giorni_decaden[val_id_slot] = val_giorni

        col_id_puro = "Identificativo" if "Identificativo" in df_schede.columns else str(df_schede.columns)
        lista_giorni_finali = []
        lista_identificativi_puliti = []
        
        for idx, riga_scheda in df_schede.iterrows():
            id_grezzo = str(riga_scheda[col_id_puro]).strip()
            if id_grezzo.endswith('.0'): id_grezzo = id_grezzo[:-2]
            if id_grezzo in ["", "nan", "None", "0", "0.0"]:
                lista_identificativi_puliti.append("")
                lista_giorni_finali.append("0")
            else:
                lista_identificativi_puliti.append(id_grezzo)
                lista_giorni_finali.append(mappa_giorni_decaden.get(id_grezzo, "0"))

        # 🛡️ 4. ASSEMBLAGGIO TABELLA REALE INTEGRATA DALLO SCREENSHOT
        df_final_view = pd.DataFrame()
        df_final_view["Identificativo"] = lista_identificativi_puliti
        df_final_view["Nome"] = df_schede["Nome"].astype(str) if "Nome" in df_schede.columns else ""
        
        col_prov_reale = "Identificativo Prov" if "Identificativo Prov" in df_schede.columns else ("Identificativo Pr" if "Identificativo Pr" in df_schede.columns else "")
        df_final_view["Identificativo Pr"] = df_schede[col_prov_reale].astype(str) if col_prov_reale else ""
        
        df_final_view["Luogo"] = df_schede[col_luogo_schede].astype(str).apply(lambda x: x.strip())
        df_final_view["Concessionario"] = df_schede["Concessionario"].astype(str) if "Concessionario" in df_schede.columns else ""
        df_final_view["Costruttore"] = df_schede["Costruttore"].astype(str) if "Costruttore" in df_schede.columns else ""
        
        # 🎯 CORREZIONE CHIRURGICA DELLA RIGA: Converte in numero intero senza far arrabbiare il compilatore
        df_final_view["Giorni Decadenza"] = pd.to_numeric(pd.Series(lista_giorni_finali), errors='coerce').fillna(0).astype(int)
        
        # 🎯 1. RIEMPIMENTO TENDINA: Se esiste il file magazzini, prende TUTTI i depositi ufficiali estratti da Sansone
        if os.path.exists(FILE_ELENCO_MAGAZZINI):
            opzioni_totali_tendina = sorted(list(set(dizionario_magazzini.values())))
            opzioni_totali_tendina = [o for o in opzioni_totali_tendina if o not in ["", "nan", "None"]]
        else:
            tutti_i_luoghi_reali = df_final_view["Luogo"].drop_duplicates().tolist()
            opzioni_totali_tendina = sorted([str(l).strip() for l in tutti_i_luoghi_reali if str(l).strip().upper() not in ["LUOGO", "MAGAZZINO SCONOSCIUTO", ""]])
        
        # 🎯 2. ANTEPRIMA DI DEFAULT: I 4 depositi indicati da Manuela
        depositi_default_manuela = [
            "WINGAMING SRL",
            "WINGAMING SRL - DEPOSITO AWP ASSEGNATE",
            "WINGAMING SRL - DEPOSITO GUASTI",
            "WINGAMING SRL - FAUSTO LAI - Sardegna"
        ]
        
        # Pre-seleziona solo quelli tra i 4 che esistono effettivamente nell'anagrafica
        default_selezionati = [d for d in depositi_default_manuela if d in opzioni_totali_tendina]
        if not default_selezionati and len(opzioni_totali_tendina) > 0:
            default_selezionati = opzioni_totali_tendina[:1]

        st.markdown("### 🎛️ Centralina di Selezione Depositi WinGaming")
        scelta_magazzini = st.multiselect(
            "Seleziona uno o più magazzini (Di default sono attivi solo i vostri quattro principali):", 
            opzioni_totali_tendina, 
            default=default_selezionati
        )
        
        # 🎯 3. FILTRAGGIO TABELLA: Mostra le schede in base ai depositi scelti ed esclude i locali installati
        if scelta_magazzini:
            scelta_pulita = [str(x).strip() for x in scelta_magazzini]
            df_filtrato = df_final_view[df_final_view["Luogo"].isin(scelta_pulita)].copy()
        else:
            # Se l'utente svuota la tendina, di paracadute mostra comunque solo i 4 default per non far apparire i locali
            df_filtrato = df_final_view[df_final_view["Luogo"].isin(depositi_default_manuela)].copy()

        # Iniezione colonna "Installare a"
        if "Concessionario" in df_filtrato.columns:
            indice_concessionario = df_filtrato.columns.get_loc("Concessionario")
            df_filtrato.insert(indice_concessionario, "Installare a", "")
        else:
            df_filtrato["Installare a"] = ""

        # Ordinamento alfabetico per Costruttore (A-Z)
        if "Costruttore" in df_filtrato.columns:
            df_filtrato = df_filtrato.sort_values(by=["Costruttore"], ascending=True).reset_index(drop=True)

        st.markdown(f"### 📊 Registro Giacenza Hub Selezionato ({len(df_filtrato)} macchine)")
        
        # REGOLA VISIVA: Dice allo schermo di trattare la colonna come numero intero (%d) eliminando i decimali spuri
        st.dataframe(
            df_filtrato, 
            hide_index=True, 
            use_container_width=True,
            column_config={
                "Giorni Decadenza": st.column_config.NumberColumn("Giorni Decadenza", format="%d")
            }
        )
        
        # Generatore Excel openpyxl
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine='openpyxl') as w: 
            df_filtrato.to_excel(w, index=False, sheet_name="Giacenza_Depositi")
            workbook = w.book
            worksheet = w.sheets["Giacenza_Depositi"]
            import openpyxl
            max_colonna_lettera = openpyxl.utils.get_column_letter(df_filtrato.shape[1])
            worksheet.auto_filter.ref = f"A1:{max_colonna_lettera}{len(df_filtrato) + 1}"
            
            bordo_sottile = Border(
                left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
                top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000')
            )
            for row in worksheet.iter_rows(min_row=1, max_row=len(df_filtrato) + 1, min_col=1, max_col=df_filtrato.shape[1]):
                for cell in row: 
                    cell.border = bordo_sottile
                    # Forza il formato numero intero nell'Excel per evitare decimali all'export
                    try:
                        if cell.column_letter == openpyxl.utils.get_column_letter(df_filtrato.columns.get_loc("Giorni Decadenza") + 1) and row.row > 1:
                            cell.number_format = '#,##0'
                    except Exception:
                        pass
            
            for col_idx in range(1, df_filtrato.shape[1] + 1):
                max_len = 0
                col_lettera = openpyxl.utils.get_column_letter(col_idx)
                for row_idx in range(1, len(df_filtrato) + 2):
                    val_cella = worksheet.cell(row=row_idx, column=col_idx).value
                    if val_cella: max_len = max(max_len, len(str(val_cella)))
                worksheet.column_dimensions[col_lettera].width = max(max_len + 4, 12)
                
        st.markdown("---")
        st.download_button(label="📥 ESPORTA TABELLA IN EXCEL (BORDATO E FORMATTATO)", data=buf.getvalue(), file_name="Giacenza_Depositi_WinGaming.xlsx")
        
    except Exception as e_grid: st.error(f"Errore caricamento griglia: {str(e_grid)}")
else:
    st.info("⏳ I file Excel stanno arrivando sul server! Clicca sul pompante in cima per generare la prima giacenza dei depositi fisici!")
